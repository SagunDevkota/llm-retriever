"""Compare RAGAS scores across RAG types, averaged per question type, with charts.

For each corpus folder (default: ./django-rest-framework and ./playwright) and each
retriever found in it:
  questions_<retriever>[_<corpus>].json -> the `category` (question type) per question,
                                           plus the `llm_cost` breakdown of the run
  ragas_results_<retriever>.xlsx        -> per-question metric scores

Retrievers whose files are missing (e.g. KETRAG has no Playwright run yet) are skipped.

Output: evaluation_report_<corpus>.xlsx in the corpus folder, with two sheets:
  RAG Comparison -> one row per (question type, RAG type), RAGs side by side, plus the
                    average retrieved-context length and the average LLM cost/token usage.
  Charts         -> a question type x RAG table per chart metric, each with a clustered
                    column chart next to it.

Cost is computed from token counts at fixed per-million prices (same rates as
cost_calculation.py), so every retriever is priced on the same basis. The provider-billed
`upstream_inference_cost` is not comparable across retrievers (the LLM retriever's routing
runs on Hetzner and bills $0, RankGPT's routing runs on OpenRouter), so it is kept only as
`billed_cost` for reference. Routing = the tree-walk / rerank calls, generation = the
final answer call; single-shot retrievers have no routing, so routing_* is 0.

Usage:
  uv run python generate_evaluation_report.py                      # both corpora
  uv run python generate_evaluation_report.py path/to/playwright   # specific folder(s)
"""

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent
DEFAULT_DIRS = [HERE / "django-rest-framework", HERE / "playwright"]

# folder name -> report file suffix.
REPORT_SUFFIX = {"django-rest-framework": "drf"}

# retriever key -> label shown in the report.
RETRIEVERS = {
    "llm_retriever": "LLM Retriever",
    "naive_rag": "Naive RAG",
    "raptor": "RAPTOR",
    "hybrid_rerank": "Hybrid + Rerank",
    "ketrag": "KETRAG",
    "rankgpt": "RankGPT",
}

# USD per 1M tokens; keep in sync with cost_calculation.py.
ROUTING_INPUT_PRICE_PER_M = 0.03
ROUTING_OUTPUT_PRICE_PER_M = 0.32
GENERATION_INPUT_PRICE_PER_M = 0.25
GENERATION_OUTPUT_PRICE_PER_M = 1.50

METRICS = [
    "context_precision",
    "context_recall",
    "helpfulness",
    "answer_accuracy",
    "factual_correctness",
    "rouge_score(mode=precision)",
    "rouge_score(mode=recall)",
]
COST_COLS = [
    "llm_calls",
    "total_cost",
    "total_tokens",
    "routing_cost",
    "routing_tokens",
    "generation_cost",
    "generation_tokens",
    "billed_cost",
]
CATEGORY_ORDER = ["straightforward", "multi_hop", "intent_driven", "debugging"]

# report column -> (chart title, y-axis title, number format, scale factor, fixed 0-100 axis).
# Score charts are put on a common 0-100 scale: 0-1 metrics x100, helpfulness (1-5) x20.
# The RAG Comparison sheet keeps the raw values.
CHARTS = {
    "context_precision": ("Context Precision by Question Type (higher is better)", "Score (0-100)", "0.0", 100, True),
    "context_recall": ("Context Recall by Question Type (higher is better)", "Score (0-100)", "0.0", 100, True),
    "answer_accuracy": ("Answer Accuracy (higher is better)", "Score (0-100)", "0.0", 100, True),
    "helpfulness": ("Helpfulness (higher is better)", "Score (0-100)", "0.0", 20, True),
    "factual_correctness": ("Factual Correctness (higher is better)", "Score (0-100)", "0.0", 100, True),
    "total_cost": ("Total Cost (lower is better)", "USD per question", "0.000000", 1, False),
    "generation_cost": ("Generation Cost (lower is better)", "USD per question", "0.000000", 1, False),
    "total_tokens": ("Total Tokens (lower is better)", "Tokens per question", "#,##0", 1, False),
}


def context_length(chunks):
    """Total character length of a question's retrieved contexts."""
    return sum(len(str(c)) for c in chunks or [])


def cost_breakdown(usage):
    """Flatten one question's `llm_cost` into the routing / generation split.

    Single-shot retrievers have no `routing` sub-dict and the whole spend is
    generation. Retrievers that skip routing (KETRAG) report its token counts as None.
    """
    routing = usage.get("routing") or {}
    generation = usage.get("generation", usage)

    r_in = routing.get("prompt_tokens") or 0
    r_out = routing.get("completion_tokens") or 0
    g_in = generation.get("prompt_tokens") or 0
    g_out = generation.get("completion_tokens") or 0

    routing_cost = (r_in * ROUTING_INPUT_PRICE_PER_M + r_out * ROUTING_OUTPUT_PRICE_PER_M) / 1e6
    generation_cost = (g_in * GENERATION_INPUT_PRICE_PER_M + g_out * GENERATION_OUTPUT_PRICE_PER_M) / 1e6
    return {
        "llm_calls": usage.get("calls"),
        "total_cost": routing_cost + generation_cost,
        "total_tokens": r_in + r_out + g_in + g_out,
        "routing_cost": routing_cost,
        "routing_tokens": r_in + r_out,
        "generation_cost": generation_cost,
        "generation_tokens": g_in + g_out,
        "billed_cost": usage.get("upstream_inference_cost"),
    }


def questions_file(data_dir, retriever):
    """questions_<retriever>.json, or the corpus-suffixed variant (questions_hybrid_rerank_playwright.json)."""
    for path in (data_dir / f"questions_{retriever}.json", data_dir / f"questions_{retriever}_{data_dir.name}.json"):
        if path.exists():
            return path
    return None


def load(data_dir, retriever):
    q_path = questions_file(data_dir, retriever)
    ragas_path = data_dir / f"ragas_results_{retriever}.xlsx"
    if q_path is None or not ragas_path.exists():
        print(f"  skipping {retriever}: missing questions or ragas results")
        return None

    questions = json.loads(q_path.read_text())
    cats = {q["user_input"]: q["category"] for q in questions}
    costs = {q["user_input"]: cost_breakdown(q["llm_cost"]) for q in questions}
    # From the JSON, not the xlsx: Excel truncates cells at 32,767 characters.
    lengths = {q["user_input"]: context_length(q.get("llm_references")) for q in questions}

    df = pd.read_excel(ragas_path)
    df["question_type"] = df["user_input"].map(cats)
    df["context_length"] = df["user_input"].map(lengths)
    for col in COST_COLS:
        df[col] = df["user_input"].map(lambda q, c=col: costs.get(q, {}).get(c, float("nan")))
    df["rag"] = RETRIEVERS[retriever]
    return df


def build_report(data_dir):
    frames = [df for r in RETRIEVERS if (df := load(data_dir, r)) is not None]
    rows = pd.concat(frames, ignore_index=True)

    report = (
        rows.groupby(["question_type", "rag"])[METRICS + ["context_length"] + COST_COLS]
        .mean()
        .reset_index()
    )
    # costs are fractions of a cent, so they need more decimals than the 0-1 metrics.
    report[METRICS + ["context_length"]] = report[METRICS + ["context_length"]].round(4)
    report[COST_COLS] = report[COST_COLS].round(8)
    report["question_type"] = pd.Categorical(report["question_type"], CATEGORY_ORDER, ordered=True)
    report["rag"] = pd.Categorical(report["rag"], list(RETRIEVERS.values()), ordered=True)
    return report.sort_values(["question_type", "rag"]).reset_index(drop=True)


def add_charts(path, report):
    """Write a question type x RAG table per chart metric, each with a column chart beside it.

    Every corpus gets the same RAG columns (all of RETRIEVERS, in order) so chart labels
    and colours line up across reports; a retriever with no run is left blank.
    """
    rag_order = list(RETRIEVERS.values())
    wb = load_workbook(path)
    ws = wb.create_sheet("Charts")
    ws["A1"] = "RAG Comparison by Question Type"
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 18
    for i in range(len(rag_order)):
        ws.column_dimensions[chr(ord("B") + i)].width = 16

    chart_col = chr(ord("B") + len(rag_order) + 1)
    block_height = 20  # rows per table + chart block
    row = 3
    for metric, (title, y_title, fmt, scale, percent_axis) in CHARTS.items():
        pivot = report.pivot(index="question_type", columns="rag", values=metric) * scale
        pivot = pivot.reindex(index=CATEGORY_ORDER, columns=rag_order)

        ws.cell(row, 1, title).font = Font(bold=True)
        header = row + 1
        ws.cell(header, 1, "Question Type").font = Font(bold=True)
        for j, rag in enumerate(rag_order, start=2):
            ws.cell(header, j, rag).font = Font(bold=True)
        for i, category in enumerate(CATEGORY_ORDER, start=header + 1):
            ws.cell(i, 1, category)
            for j, rag in enumerate(rag_order, start=2):
                value = pivot.loc[category, rag]
                cell = ws.cell(i, j, None if pd.isna(value) else float(value))
                cell.number_format = fmt
        last = header + len(CATEGORY_ORDER)

        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        chart.title = title
        chart.y_axis.title = y_title
        chart.x_axis.title = "Question Type"
        chart.y_axis.number_format = fmt
        if percent_axis:
            chart.y_axis.scaling.min = 0
            chart.y_axis.scaling.max = 100
            chart.y_axis.majorUnit = 20
        # openpyxl hides both axes by default.
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.legend.position = "b"
        # openpyxl leaves `overlay` unset, which Excel treats as "draw on top of the plot";
        # turn it off and pin the plot area so title, axis titles and legend get their own space.
        chart.title.overlay = False
        chart.x_axis.title.overlay = False
        chart.y_axis.title.overlay = False
        chart.legend.overlay = False
        chart.layout = Layout(
            manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge", x=0.09, y=0.12, w=0.88, h=0.62)
        )
        data = Reference(ws, min_col=2, max_col=1 + len(rag_order), min_row=header, max_row=last)
        cats = Reference(ws, min_col=1, min_row=header + 1, max_row=last)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 22
        chart.height = 9.5
        ws.add_chart(chart, f"{chart_col}{row}")

        row += block_height

    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("data_dirs", nargs="*", type=Path, default=DEFAULT_DIRS, help="corpus folders")
    args = parser.parse_args()

    for data_dir in args.data_dirs:
        data_dir = data_dir.resolve()
        print(f"\n=== {data_dir.name} ===")
        report = build_report(data_dir)

        output = data_dir / f"evaluation_report_{REPORT_SUFFIX.get(data_dir.name, data_dir.name)}.xlsx"
        report.to_excel(output, index=False, sheet_name="RAG Comparison")
        add_charts(output, report)

        print(report.to_string(index=False))
        print(f"\nWrote report -> {output}")


if __name__ == "__main__":
    main()
